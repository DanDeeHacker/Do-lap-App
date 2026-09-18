"""Generate a daily subjective-log template (.xlsx + .csv) designed to feed
backtest_engine.py months later. Keyed by absolute date, structured scales,
dropdowns, and a separate injury-episode log (the ground-truth outcome).

Design rules baked into the layout (see the 'Návod' sheet):
  1. Log BLIND — before opening the app's score.
  2. One row per calendar day, rest days included ("OK today" is data).
  3. Log within ~24h, never reconstruct from memory.
  4. Injury onset date is the sacred label — everything is validated against it.

Usage:  python make_daily_log_template.py [days] [out.xlsx]
"""
import csv
import sys
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DAYS = int(sys.argv[1]) if len(sys.argv) > 1 else 180
OUT = sys.argv[2] if len(sys.argv) > 2 else "/Users/danieltrnovec/Downloads/dosslap_daily_log_template.xlsx"
CSV_OUT = OUT.rsplit(".", 1)[0] + "_denni.csv"

DOW = ["Po", "Út", "St", "Čt", "Pá", "So", "Ne"]
SITES = ('"Achillova šlacha,lýtko,holeň,koleno,hamstring,kvadriceps,kyčel,'
         'hýždě,chodidlo,plantární fascie,kotník,IT pás,tříslo,bederní páteř,jiné"')
HEAD_FILL = PatternFill("solid", fgColor="0A2540")
NOTE_FILL = PatternFill("solid", fgColor="EAF1F8")


def _style_header(ws, cols, row=1):
    for i, (_, _, w) in enumerate(cols, 1):
        c = ws.cell(row=row, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"B{row + 1}"     # string form — ws.cell() would instantiate a blank row
    ws.row_dimensions[row].height = 30


def _dv_list(ws, formula, col, n_rows):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{n_rows + 1}")


def _dv_whole(ws, lo, hi, col, n_rows):
    dv = DataValidation(type="whole", operator="between", formula1=str(lo), formula2=str(hi), allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{n_rows + 1}")


def build():
    wb = Workbook()

    # ---------------------------------------------------------------- Denní
    ws = wb.active
    ws.title = "Denní"
    cols = [
        ("datum", "date", 12), ("den", "dow", 6),
        ("ok_dnes", "Nic mě netrápí (A/N)", 10),
        ("nohy_1_5", "Nohy 1=těžké · 5=svěží", 12),
        ("bolest_0_10", "Bolest 0–10", 11),
        ("misto", "Kde bolí", 18), ("strana", "Strana", 8), ("typ", "Charakter", 12),
        ("svaly_0_10", "Svalovka 0–10", 11),
        ("spanek_1_5", "Kvalita spánku 1–5", 12),
        ("stres_1_5", "Životní stres 1–5", 11),
        ("konfoundery", "Konfoundery (nemoc/alkohol/cesta/nové boty/nový povrch)", 30),
        ("pozn", "Poznámka", 34),
    ]
    ws.append([c[1] for c in cols])
    _style_header(ws, cols)
    d0 = date.today()
    for i in range(DAYS):
        d = d0 + timedelta(days=i)
        ws.append([d.isoformat(), DOW[d.weekday()], "", "", "", "", "", "", "", "", "", "", ""])
    n = DAYS
    _dv_list(ws, '"A,N"', "C", n)
    _dv_whole(ws, 1, 5, "D", n)
    _dv_whole(ws, 0, 10, "E", n)
    _dv_list(ws, SITES, "F", n)
    _dv_list(ws, '"L,P,obě,—"', "G", n)
    _dv_list(ws, '"ostrá,tupá,pálivá,ztuhlost,jiné"', "H", n)
    _dv_whole(ws, 0, 10, "I", n)
    _dv_whole(ws, 1, 5, "J", n)
    _dv_whole(ws, 1, 5, "K", n)
    for r in range(2, n + 2):                       # tint the note column
        ws.cell(row=r, column=len(cols)).fill = NOTE_FILL

    # ---------------------------------------------------------------- Běhy
    wr = wb.create_sheet("Běhy")
    rcols = [
        ("datum", "date", 12), ("beh", "Běh / typ", 20),
        ("rpe_1_10", "RPE 1–10", 10), ("pocit_1_5", "Pocit 1–5", 10),
        ("nohy_1_5", "Nohy 1–5", 10), ("ztuhlost_pred_1_5", "Ztuhlost před 1–5", 13),
        ("bolest_beh_0_10", "Bolest při běhu 0–10", 13),
        ("misto", "Kde", 18), ("strana", "Strana", 8),
        ("niggle", "Niggle (A/N)", 10), ("pozn", "Poznámka", 34),
    ]
    wr.append([c[1] for c in rcols])
    _style_header(wr, rcols)
    RN = 200
    _dv_whole(wr, 1, 10, "C", RN)
    _dv_whole(wr, 1, 5, "D", RN)
    _dv_whole(wr, 1, 5, "E", RN)
    _dv_whole(wr, 1, 5, "F", RN)
    _dv_whole(wr, 0, 10, "G", RN)
    _dv_list(wr, SITES, "H", RN)
    _dv_list(wr, '"L,P,obě,—"', "I", RN)
    _dv_list(wr, '"A,N"', "J", RN)

    # ---------------------------------------------------- Zranění (epizody)
    wz = wb.create_sheet("Zranění (epizody)")
    zcols = [
        ("onset_datum", "Začátek (onset) — SACRED", 18),
        ("misto", "Kde", 18), ("strana", "Strana", 8), ("tkan_odhad", "Tkáň (odhad)", 16),
        ("mechanismus", "Co to spustilo", 30),
        ("q_ucast", "OSTRC účast 0/8/17/25", 12), ("q_objem", "OSTRC objem", 11),
        ("q_vykon", "OSTRC výkon", 11), ("q_bolest", "OSTRC bolest", 11),
        ("dny_omezeno", "Dní omezeno", 11), ("dny_vynechano", "Dní vynecháno", 12),
        ("konec_datum", "Vyřešeno (datum)", 15), ("pozn", "Poznámka", 34),
    ]
    wz.append([c[1] for c in zcols])
    _style_header(wz, zcols)
    ZN = 60
    _dv_list(wz, SITES, "B", ZN)
    _dv_list(wz, '"L,P,obě,—"', "C", ZN)
    for col in ("F", "G", "H", "I"):
        _dv_list(wz, '"0,8,17,25"', col, ZN)

    # ---------------------------------------------------- Týdenní OSTRC
    wo = wb.create_sheet("Týdenní OSTRC")
    ocols = [
        ("tyden_od", "Týden od (datum)", 15),
        ("q_ucast", "Účast 0/8/17/25", 12), ("q_objem", "Objem", 10),
        ("q_vykon", "Výkon", 10), ("q_bolest", "Bolest", 10),
        ("severita", "Severita (=součet, 0–100)", 16), ("pozn", "Poznámka", 34),
    ]
    wo.append([c[1] for c in ocols])
    _style_header(wo, ocols)
    ON = 40
    for col in ("B", "C", "D", "E"):
        _dv_list(wo, '"0,8,17,25"', col, ON)
    for r in range(2, ON + 2):                      # auto severity = sum of the four
        wo.cell(row=r, column=6, value=f"=SUM(B{r}:E{r})")

    # ---------------------------------------------------------------- Návod
    wn = wb.create_sheet("Návod")
    for row in [
        ["Deník subjektivních dat — návod", ""],
        ["", ""],
        ["4 pravidla, na kterých stojí pozdější backtest", ""],
        ["1. NASLEPO", "Vyplň dřív, než otevřeš skóre v aplikaci — jinak hodnotíš engine, ne své tělo."],
        ["2. Každý den", "Jeden řádek na kalendářní den, i dny volna. Nic mě netrápí = A je taky data."],
        ["3. Do 24 h", "Nevyplňuj zpětně z paměti víc než den — je to zdroj chyb v datech onsetu."],
        ["4. Onset je svatý", "Zranění zapiš dnem, kdy poprvé omezilo běh (ne až když bylo zlé). Podle toho se vše validuje."],
        ["", ""],
        ["Listy", ""],
        ["Denní", "30 s ráno naslepo. Většinu dní vyplníš jen ok_dnes=A, nohy, spánek, stres; bolest 0."],
        ["Běhy", "20 s po každém běhu (nebo to zadej rovnou v aplikaci na obrazovce hodnocení)."],
        ["Zranění (epizody)", "Jedna řádka na zranění. Onset zapiš hned; dny/konec doplň průběžně."],
        ["Týdenní OSTRC", "1× týdně validovaný dotazník přetížení — páteřní outcome i bez konkrétního zranění."],
        ["", ""],
        ["Škály", ""],
        ["Nohy 1–5", "1 = těžké/unavené · 5 = svěží"],
        ["Bolest / Svalovka 0–10", "0 = žádná · 10 = musím přerušit"],
        ["Spánek / Stres 1–5", "spánek 1 = mizerný, 5 = skvělý · stres 1 = klid, 5 = vysoký"],
        ["OSTRC 0/8/17/25", "0 = bez potíží · 8 = mírné · 17 = střední/omezení · 25 = velké/nemohl jsem"],
        ["Konfoundery", "nemoc, alkohol, cesta/časový posun, špatné prostředí na spánek, nové boty, nový povrch — "
                        "HRV a klidový tep na ně reagují stejně jako na trénink, proto je nutné je odlišit."],
        ["", ""],
        ["K čemu to je", "Za pár měsíců to načtu do backtest_engine.py a spočítám, jestli skóre/kvadrant "
                        "stoupl PŘED onsetem zranění a s jakým předstihem (lead-time). Dvě klíčová pole: "
                        "datum onsetu a denní bolest/svalovka zapsaná naslepo."],
    ]:
        wn.append(row)
    wn["A1"].font = Font(bold=True, size=13)
    for r in (3,):
        wn.cell(row=r, column=1).font = Font(bold=True)
    wn.column_dimensions["A"].width = 26
    wn.column_dimensions["B"].width = 100
    # move Návod to the front for visibility
    wb.move_sheet("Návod", -(len(wb.sheetnames) - 1))

    wb.save(OUT)

    # plain CSV of the Denní sheet (headers + prefilled dates)
    with open(CSV_OUT, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow([c[0] for c in cols])
        for i in range(DAYS):
            d = d0 + timedelta(days=i)
            w.writerow([d.isoformat(), DOW[d.weekday()], "", "", "", "", "", "", "", "", "", "", ""])

    print(f"Hotovo:\n  {OUT}\n  {CSV_OUT}\n  {DAYS} dní od {d0.isoformat()} do {(d0+timedelta(days=DAYS-1)).isoformat()}")


if __name__ == "__main__":
    build()
