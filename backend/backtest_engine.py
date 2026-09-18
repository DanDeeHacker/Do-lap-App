"""Historical replay of injury-risk engine v0.4 over one runner's real data.

For each weekly as-of date it rebuilds an in-memory DB holding only the
activities/daily metrics that existed *up to that date*, pins the engine's
notion of "today" to that date, runs assess(), and records the resulting
quadrant / tier / scores / driving signals. Writes an .xlsx timeline so you
can see when — and why — the quadrant would have flipped.

Usage:  python backtest_engine.py [runner_id] [out.xlsx]
"""
import sys
from datetime import date, timedelta

from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.db import Base
from app import models
from app.metrics import engine as E

RID = sys.argv[1] if len(sys.argv) > 1 else "run-0012"
OUT = sys.argv[2] if len(sys.argv) > 2 else "/Users/danieltrnovec/Downloads/dosslap_engine_v04_backtest.xlsx"
DB = "dosslap.db"

QUAD_CZ = {"stable": "Stabilní", "overreaching": "Přetížení", "silent": "Tichý drift", "critical": "Kritické"}
QUAD_FILL = {"stable": "C6EFCE", "overreaching": "FFEB9C", "silent": "BDD7EE", "critical": "FFC7CE"}
TIER_CZ = {"ok": "OK", "watch": "Sledovat", "alert": "Alarm"}


def _rows_of(session, model, rid):
    cols = [c.name for c in model.__table__.columns]
    out = []
    for r in session.query(model).filter(model.runner_id == rid).all():
        out.append({c: getattr(r, c) for c in cols})
    return out


def main():
    src = sessionmaker(bind=create_engine(f"sqlite:///{DB}"))()
    runner = src.query(models.Runner).filter(models.Runner.id == RID).first()
    if not runner:
        print(f"Runner {RID} nenalezen v {DB}")
        return
    runner_data = {c.name: getattr(runner, c.name) for c in models.Runner.__table__.columns}
    acts = _rows_of(src, models.Activity, RID)
    daily = _rows_of(src, models.DailyMetric, RID)
    src.close()
    if not acts:
        print("Žádné aktivity.")
        return

    first = min(a["started_at"][:10] for a in acts)
    last = max(a["started_at"][:10] for a in acts)
    asof = date.fromisoformat(first) + timedelta(days=42)   # start once a baseline can plausibly exist
    end = date.fromisoformat(last)
    asofs = []
    while asof <= end:
        asofs.append(asof)
        asof += timedelta(days=7)
    if not asofs or asofs[-1] != end:
        asofs.append(end)

    orig_today = E.today_date
    results = []
    try:
        for ad in asofs:
            E.today_date = (lambda d: (lambda: d))(ad)   # pin engine "today" to as-of
            eng = create_engine("sqlite://")
            Base.metadata.create_all(eng)
            ts = sessionmaker(bind=eng)()
            ts.add(models.Runner(**runner_data))
            cut = ad.isoformat()
            for a in acts:
                if a["started_at"][:10] <= cut:
                    ts.add(models.Activity(**{k: v for k, v in a.items() if k != "id"}))
            for d in daily:
                if d["date"][:10] <= cut:
                    ts.add(models.DailyMetric(**{k: v for k, v in d.items() if k != "id"}))
            ts.commit()

            # seed the previous week's quadrant so the engine's hysteresis
            # (quadrant_of) sees prior state, exactly as it would in production
            if results:
                ts.add(models.Assessment(
                    runner_id=RID, quadrant=results[-1]["_quad_key"], tier="ok",
                    mech=0, load=0, symp=0, overall=0, engine_version="v0.4",
                ))
                ts.commit()

            a = E.assess(ts, RID)
            conf = a["confidence"] or {}
            L = a.get("loadDetail") or {}
            tv, gc, rcv = a.get("tavr"), a.get("gct"), a.get("rcv")
            sigs = a.get("signals") or []
            n_to_date = sum(1 for x in acts if x["started_at"][:10] <= cut)
            results.append({
                "as_of": cut,
                "runs_to_date": n_to_date,
                "days_history": conf.get("days"),
                "baseline_sessions": conf.get("baseSessions"),
                "matched_sessions": conf.get("sessions"),
                "confidence": conf.get("value"),
                "acute_km": L.get("acute"), "chronic_km": L.get("chronic"),
                "ewma_ratio": L.get("ratio") if L.get("valid") else None,
                "monotony": L.get("monotony"),
                "tavr_z": (tv or {}).get("z"), "gct_z": (gc or {}).get("z"),
                "hrv_z": ((rcv or {}).get("hrv") or {}).get("z"),
                "rhr_z": ((rcv or {}).get("rhr") or {}).get("z"),
                "mech": a["mech"], "load": a["load"], "symp": a["symp"], "overall": a["overall"],
                "tier": TIER_CZ.get(a["tier"], a["tier"]),
                "quadrant": QUAD_CZ.get(a["quadrant"], a["quadrant"]),
                "_quad_key": a["quadrant"],
                "top_signals": " · ".join(f"{s['name']} (+{s['pts']}, {s['grade']})" for s in sigs[:3]) or "—",
            })
            ts.close()
            eng.dispose()
    finally:
        E.today_date = orig_today

    # mark quadrant transitions
    prev = None
    for r in results:
        r["quad_changed"] = "← ZMĚNA" if (prev is not None and r["_quad_key"] != prev) else ""
        prev = r["_quad_key"]

    _write_xlsx(results, runner_data, first, last)
    print(f"Hotovo: {len(results)} týdenních snímků → {OUT}")
    changes = [r for r in results if r["quad_changed"]]
    print(f"Změn kvadrantu: {len(changes)}")
    for r in changes:
        print(f"  {r['as_of']}: → {r['quadrant']} (overall {r['overall']}, mech {r['mech']}, load {r['load']})")


def _write_xlsx(results, runner_data, first, last):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Engine v0.4 backtest"

    title = [
        [f"Engine v0.4 — historický replay", runner_data.get("name", "")],
        ["Období dat", f"{first} → {last}"],
        ["Kvadrant = mechanika × zátěž", "Stabilní / Přetížení (zátěž↑) / Tichý drift (mechanika driftuje bez objemu) / Kritické (obojí)"],
        ["Pozn.", "Každý řádek = stav enginu, kdyby se počítal jen z dat do daného data. z = odchylka od vlastní normy."],
        [],
    ]
    for row in title:
        ws.append(row)
    ws["A1"].font = Font(bold=True, size=13)

    cols = [
        ("as_of", "Datum", 12), ("runs_to_date", "Běhů celkem", 11),
        ("days_history", "Dnů historie", 12), ("baseline_sessions", "Baseline běhů", 13),
        ("confidence", "Konfidence", 11),
        ("acute_km", "Akut km/7d", 11), ("chronic_km", "Chron km", 10),
        ("ewma_ratio", "EWMA poměr", 11), ("monotony", "Monotón.", 10),
        ("tavr_z", "Vert.poměr z", 12), ("gct_z", "GCT z", 8),
        ("hrv_z", "HRV z", 8), ("rhr_z", "Klid.tep z", 10),
        ("mech", "Mechanika", 10), ("load", "Zátěž", 8), ("symp", "Příznaky", 9),
        ("overall", "Celkem", 8), ("tier", "Úroveň", 10), ("quadrant", "Kvadrant", 13),
        ("quad_changed", "Přechod", 10), ("top_signals", "Hlavní signály", 60),
    ]
    header_row = ws.max_row + 1
    ws.append([c[1] for c in cols])
    hfill = PatternFill("solid", fgColor="0A2540")
    for i, (_, _, w) in enumerate(cols, 1):
        cell = ws.cell(row=header_row, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hfill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w

    for r in results:
        vals = []
        for key, _, _ in cols:
            v = r.get(key)
            if isinstance(v, float):
                v = round(v, 2)
            vals.append(v)
        ws.append(vals)
        row_i = ws.max_row
        qcell = ws.cell(row=row_i, column=[c[0] for c in cols].index("quadrant") + 1)
        qcell.fill = PatternFill("solid", fgColor=QUAD_FILL.get(r["_quad_key"], "FFFFFF"))
        if r["quad_changed"]:
            ws.cell(row=row_i, column=[c[0] for c in cols].index("quad_changed") + 1).font = Font(bold=True, color="C00000")

    ws.freeze_panes = ws.cell(row=header_row + 1, column=2)

    # legend sheet
    ws2 = wb.create_sheet("Legenda")
    for row in [
        ["Sloupec", "Význam"],
        ["Konfidence", "Spolehlivost baseline 0–1. Pod 0,6 se mechanické signály (vert.poměr, GCT) vůbec nepočítají."],
        ["EWMA poměr", "Akutní vs chronická zátěž. Pásmo ~0,7–1,3 klidné; >1,5 zřetelně zvýšené riziko; počítá se od chron. ≥15 km."],
        ["Vert.poměr z / GCT z", "Odchylka běžecké mechaniky od vlastní normy ve srovnatelném terénu/tempu. ≥1 = drift."],
        ["HRV z / Klid.tep z", "Regenerace z hodinek proti baseline. HRV z ≤ −1 nebo klid.tep z ≥ 1,2 = potlačená regenerace."],
        ["Mechanika / Zátěž / Příznaky", "Tři osy skóre 0–100, nesčítají se do jednoho. Kvadrant je jejich kombinace."],
        ["Kvadrant", "Stabilní | Přetížení (zátěž↑, technika drží) | Tichý drift (technika driftuje bez objemu) | Kritické (obojí)"],
        ["Přechod", "Řádek, kde se kvadrant změnil proti předchozímu týdnu — moment, který stojí za pohled."],
    ]:
        ws2.append(row)
    ws2["A1"].font = Font(bold=True)
    ws2.column_dimensions["A"].width = 26
    ws2.column_dimensions["B"].width = 110

    wb.save(OUT)


if __name__ == "__main__":
    main()
