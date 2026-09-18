"""Backtest + live-formula workbook for engine v0.6.0.

v0.6 rebuilds the load axis around the single-session paradigm (BJSM 2025):
single-session distance spike + 28-day latent memory + pace spike become the
spine, ACWR is demoted to grade-C context, and a load×capacity (readiness)
interaction and injury-history frailty multiplier are added. This workbook
mirrors that exact chain as editable Excel formulas next to the engine's numbers.

Sheets:
  1. Kalkulace (živé vzorce) : inputs → per-signal points → axes → overall → quadrant.
  2. Backtest — týdně / 3. denně : assess() replay over history.
  4. Signály.

Injury/outcome labelling for validation follows the Yamato 2015 consensus RRI
definition (pain restricting running ≥7 days or 3 sessions, or requiring a
professional consult) — the standard the backtest's positive cases use.
Historical replay is the objective wearable engine (self-report not replayed),
so symptom-axis self-report signals (pain/complaints) and the injury-history
frailty are inactive in the timeline unless present on the Runner record.
"""
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app import models
from app.db import Base
from app.metrics import engine as E
from app.metrics import sig_doc as SD

DB = "dosslap.db"
RID = "run-0012"
OUT = "/Users/danieltrnovec/Downloads/dosslap/dosslap_engine_v060_backtest.xlsx"
QUADL = {k: v["t"] for k, v in SD.QUAD.items()}

HEAD = PatternFill("solid", fgColor="0A2540")
GRP = PatternFill("solid", fgColor="1B4A6B")
QFILL = {"stable": "C6EFCE", "overreaching": "FFEB9C", "silent": "BDD7EE", "critical": "FFC7CE"}
IN_FILL = PatternFill("solid", fgColor="EAF1F8")
F_FILL = PatternFill("solid", fgColor="FFF6E5")
RES_FILL = PatternFill("solid", fgColor="E7F5EC")


def header(ws, names, widths, row=1, fill=HEAD):
    for i, n in enumerate(names, 1):
        c = ws.cell(row=row, column=i, value=n)
        c.font = Font(bold=True, color="FFFFFF" if fill in (HEAD, GRP) else "000000")
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.row_dimensions[row].height = 28


def _rows_of(session, model, rid):
    cols = [c.name for c in model.__table__.columns]
    return [{c: getattr(r, c) for c in cols} for r in session.query(model).filter(model.runner_id == rid).all()]


def replay(step_days=7):
    src = sessionmaker(bind=create_engine(f"sqlite:///{DB}"))()
    runner = src.query(models.Runner).filter(models.Runner.id == RID).first()
    rdata = {c.name: getattr(runner, c.name) for c in models.Runner.__table__.columns}
    acts = _rows_of(src, models.Activity, RID)
    daily = _rows_of(src, models.DailyMetric, RID)
    devices = _rows_of(src, models.DeviceHistory, RID)
    src.close()
    first = min(a["started_at"][:10] for a in acts)
    last = max(a["started_at"][:10] for a in acts)
    ad = date.fromisoformat(first) + timedelta(days=42)
    end = date.fromisoformat(last)
    asofs = []
    while ad <= end:
        asofs.append(ad)
        ad += timedelta(days=step_days)
    if not asofs or asofs[-1] != end:
        asofs.append(end)

    out = []
    prev_q = None
    for ad in asofs:
        with E.today_pinned(ad):  # thread-safe clock override (v0.6)
            eng = create_engine("sqlite://")
            Base.metadata.create_all(eng)
            ts = sessionmaker(bind=eng)()
            ts.add(models.Runner(**rdata))
            cut = ad.isoformat()
            for a in acts:
                if a["started_at"][:10] <= cut:
                    ts.add(models.Activity(**{k: v for k, v in a.items() if k != "id"}))
            for d in daily:
                if d["date"][:10] <= cut:
                    ts.add(models.DailyMetric(**{k: v for k, v in d.items() if k != "id"}))
            for dv in devices:
                if (dv.get("recorded_at") or "")[:10] <= cut:
                    ts.add(models.DeviceHistory(**{k: v for k, v in dv.items() if k != "id"}))
            ts.commit()
            if prev_q:
                ts.add(models.Assessment(runner_id=RID, quadrant=prev_q, tier="ok",
                                         mech=0, load=0, symp=0, overall=0, engine_version=E.ENGINE_VERSION))
                ts.commit()
            av = E.assess(ts, RID)
            L = av.get("loadDetail") or {}
            gd = av.get("gradientDescent") or {}
            rcv = av.get("rcv") or {}
            sub = lambda k, s="z": (av.get(k) or {}).get(s)
            out.append({
                "date": cut, "valid": 1 if L.get("valid") else 0,
                "tavr_z": sub("tavr"), "gct_z": sub("gct"), "cad_z": sub("cadence"), "vosc_z": sub("vosc"),
                "bal_exc": (av.get("bal") or {}).get("excursion"), "dec_trend": (av.get("dec") or {}).get("trend"),
                "gaitcv_ratio": (av.get("gaitCv") or {}).get("ratio"),
                # v0.6 load spine
                "sessionSpike": L.get("sessionSpike"), "spikeLatent": L.get("spikeLatent"),
                "paceSpike": L.get("paceSpike"), "capacityDeficit": av.get("capacityDeficit"),
                "frailty": av.get("frailty") or 1,
                "ratio": L.get("ratio"), "hiAcute": L.get("hiAcute"), "hiRatio": L.get("hiRatio"),
                "loadCreep": L.get("loadCreep"), "monotony": L.get("monotony"),
                "descentSpike": L.get("descentSpike"), "steepSpike": gd.get("steepSpike"),
                "aer_mean": (av.get("aer") or {}).get("mean"),
                "hrv_z": (rcv.get("hrv") or {}).get("z"), "rhr_z": (rcv.get("rhr") or {}).get("z"),
                "hrvcv_ratio": (av.get("hrvCv") or {}).get("ratio"),
                "tsbBalance": L.get("tsbBalance"), "chronic": L.get("chronic"),
                "sleep_debt": (rcv.get("sleep") or {}).get("debt"),
                "sleepreg_ratio": (av.get("sleepReg") or {}).get("ratio"),
                "sleepeff_now": (av.get("sleepEff") or {}).get("now"),
                # backtest extras
                "acute": L.get("acute"), "runKm7": L.get("runKm7"), "crossLoad7": L.get("crossLoad7"),
                "sessionSpikeKm": L.get("sessionSpikeKm"), "safeLongRunKm": L.get("safeLongRunKm"),
                "duty_z": (av.get("duty") or {}).get("z"),
                "regen": rcv.get("score"), "regen_d": rcv.get("scoreDelta"),
                "conf": (av.get("confidence") or {}).get("value"),
                "days": (av.get("confidence") or {}).get("days"),
                "runs": sum(1 for x in acts if x["started_at"][:10] <= cut),
                # engine reference
                "eng_mech": av["mech"], "eng_load": av["load"], "eng_symp": av["symp"],
                "eng_overall": av["overall"], "eng_quad": QUADL[av["quadrant"]], "_q": av["quadrant"],
                "signals": " · ".join(f"{s['name']} (+{s['pts']},{s['grade']})" for s in (av.get("signals") or [])[:4]) or "—",
            })
            prev_q = av["quadrant"]
            ts.close()
            eng.dispose()
    for i, r in enumerate(out):
        r["trans"] = "← ZMĚNA" if i and r["_q"] != out[i - 1]["_q"] else None
    return out, runner.name, f"{first} → {last}"


# ---- Python mirror of the v0.6 scoring, to VALIDATE the Excel formulas match engine
def _calc(r, prev_q):
    g = lambda k: r.get(k) or 0
    cl = lambda x, lo, hi: max(lo, min(hi, x))
    R = round
    v = r["valid"]
    fr = r.get("frailty") or 1
    mech = [
        R(cl(g("tavr_z") - 0.2, 0, 4) * 17), R(cl(g("gct_z") - 0.2, 0, 4) * 13),
        R(cl(-g("cad_z") - 0.2, 0, 4) * 10), R(cl(g("vosc_z") - 0.2, 0, 4) * 10),
        R(cl(g("bal_exc") - 0.4, 0, 3) * 22),
        R(cl(g("dec_trend") - 0.1, 0, 1.2) * 26) if g("dec_trend") > 0.15 else 0,
        R(cl((g("gaitcv_ratio") - 1.5) * 14, 0, 14)) if g("gaitcv_ratio") >= 1.5 else 0,
    ]
    rt = r.get("ratio")
    s = r.get("sessionSpike")
    sspike = 0
    if v and s and s > 1.1:
        if s > 2.0:
            sspike = R(cl((s - 2.0) * 16, 0, 16) + 14)
        elif s > 1.3:
            sspike = R(cl((s - 1.3) * 20, 0, 14))
        else:
            sspike = R(cl((s - 1.1) * 25, 0, 6))
    lat = r.get("spikeLatent")
    slat = R(cl(lat * 22, 0, 16)) if (v and lat) else 0
    ps = r.get("paceSpike")
    pspk = R(cl((ps - 1.06) * 40, 0, 10)) if (v and ps and ps > 1.06) else 0
    ewma = 0
    if v and rt is not None:
        if rt > 1.5:
            ewma = R(cl((rt - 1.5) * 18, 0, 12) + 2)
        elif rt < 0.7:
            ewma = 10
    cd = r.get("capacityDeficit")
    spike_sev = max((s or 1) - 1, (rt or 1) - 1, 0)
    lcap = R(cl(cd * spike_sev * 34, 0, 16)) if (v and cd is not None and cd >= 0.3 and spike_sev > 0.1) else 0
    tsb = 0
    if v and r.get("chronic") and r.get("tsbBalance") is not None:
        rel = r["tsbBalance"] / r["chronic"]
        if rel <= -0.12:
            tsb = R(cl((-rel - 0.12) * 90, 0, 12))
    load = [
        ewma, sspike, slat, pspk,
        R(cl((g("hiRatio") - 1.5) * 20, 0, 20)) if (v and g("hiAcute") >= 60 and g("hiRatio") > 1.5) else 0,
        R(cl((g("loadCreep") - 1.15) * 30, 0, 10)) if (v and rt is not None and rt < 1.3 and g("loadCreep") >= 1.15) else 0,
        R(cl((g("monotony") - 2.4) * 7, 0, 20)) if g("monotony") > 2.4 else 0,
        R(cl((g("descentSpike") - 1.45) * 15, 0, 14)) if g("descentSpike") > 1.45 else 0,
        R(cl((g("steepSpike") - 1.5) * 12, 0, 12)) if g("steepSpike") > 1.5 else 0,
        R(cl((g("aer_mean") - 5.5) * 3, 0, 12)) if g("aer_mean") > 5.5 else 0,
        R(cl(-g("hrv_z") * 10, 0, 22)) if g("hrv_z") <= -1.0 else 0,
        R(cl(g("rhr_z") * 8, 0, 18)) if g("rhr_z") >= 1.2 else 0,
        R(cl((g("hrvcv_ratio") - 1.4) * 14, 0, 10)) if g("hrvcv_ratio") >= 1.4 else 0,
        lcap, tsb,
    ]
    symp = [
        R(cl(g("sleep_debt") * 2.5, 0, 16)) if g("sleep_debt") >= 4 else 0,
        R(cl((g("sleepreg_ratio") - 1.5) * 12, 0, 14)) if g("sleepreg_ratio") >= 1.5 else 0,
        R(cl((0.85 - r["sleepeff_now"]) * 60, 0, 16)) if (r.get("sleepeff_now") and r["sleepeff_now"] < 0.85) else 0,
    ]
    M = min(100, R(sum(mech) * fr))
    Lo = min(100, R(sum(load) * fr))
    Sy = min(100, R(sum(symp)))
    ov = R(cl(M * 0.38 + Lo * 0.30 + Sy * 0.52, 0, 100))
    lh = Lo >= 25 or (prev_q in ("overreaching", "critical") and Lo >= 18)
    mh = M >= 25 or (prev_q in ("silent", "critical") and M >= 18)
    q = "critical" if (lh and mh) else "overreaching" if lh else "silent" if mh else "stable"
    return M, Lo, Sy, ov, q


# ---- Kalkulace sheet column spec: (key, header, width, kind)
KCOLS = [
    ("date", "datum", 11, "in"), ("valid", "valid", 6, "in"),
    ("tavr_z", "vert.p. z", 8, "in"), ("gct_z", "GCT z", 8, "in"), ("cad_z", "kad. z", 8, "in"),
    ("vosc_z", "osc. z", 8, "in"), ("bal_exc", "bal p.b.", 8, "in"), ("dec_trend", "dec tr.", 8, "in"),
    ("gaitcv_ratio", "gaitCV×", 8, "in"),
    ("sessionSpike", "běh spike×", 9, "in"), ("spikeLatent", "doznívá", 8, "in"), ("paceSpike", "tempo×", 8, "in"),
    ("capacityDeficit", "kapac.def", 9, "in"), ("frailty", "frailty×", 8, "in"),
    ("ratio", "ACWR", 7, "in"), ("hiAcute", "HI akut", 8, "in"), ("hiRatio", "HI ×", 7, "in"),
    ("loadCreep", "creep×", 8, "in"), ("monotony", "mono", 7, "in"), ("descentSpike", "sběh ×", 8, "in"),
    ("steepSpike", "strmé ×", 8, "in"), ("aer_mean", "aerob %", 8, "in"), ("hrv_z", "HRV z", 7, "in"),
    ("rhr_z", "tep z", 7, "in"), ("hrvcv_ratio", "HRV CV×", 8, "in"), ("tsbBalance", "TSB", 7, "in"),
    ("chronic", "chron.", 8, "in"),
    ("sleep_debt", "sp. dluh", 8, "in"), ("sleepreg_ratio", "sp. SD×", 8, "in"), ("sleepeff_now", "sp. efek", 8, "in"),
    ("p_tavr", "+vert.p.", 7, "f"), ("p_gct", "+GCT", 7, "f"), ("p_cad", "+kad.", 7, "f"), ("p_vosc", "+osc.", 7, "f"),
    ("p_bal", "+bal.", 7, "f"), ("p_dec", "+dec", 7, "f"), ("p_gaitcv", "+gaitCV", 7, "f"),
    ("p_ewma", "+ACWR", 7, "f"), ("p_sspike", "+běh sp.", 8, "f"), ("p_slatent", "+doznívá", 8, "f"),
    ("p_pspike", "+tempo", 7, "f"), ("p_hi", "+HI", 7, "f"), ("p_creep", "+creep", 7, "f"), ("p_mono", "+mono", 7, "f"),
    ("p_desc", "+sběh", 7, "f"), ("p_steep", "+strmé", 7, "f"), ("p_aer", "+aerob", 7, "f"), ("p_hrv", "+HRV", 7, "f"),
    ("p_rhr", "+tep", 7, "f"), ("p_hrvcv", "+HRVcv", 7, "f"), ("p_lcap", "+kap×zát", 8, "f"), ("p_tsb", "+TSB", 7, "f"),
    ("p_sleep", "+spánek", 7, "f"), ("p_sreg", "+sp.SD", 7, "f"), ("p_seff", "+sp.ef", 7, "f"),
    ("MECH", "MECH", 8, "f"), ("LOAD", "ZÁTĚŽ", 8, "f"), ("SYMP", "PŘÍZN.", 8, "f"),
    ("OVERALL", "CELKEM", 8, "f"), ("TIER", "úroveň", 9, "f"), ("QUAD", "KVADRANT", 16, "f"),
    ("eng_mech", "e:mech", 7, "in"), ("eng_load", "e:zát", 7, "in"), ("eng_symp", "e:přízn", 7, "in"),
    ("eng_overall", "e:celk", 7, "in"), ("eng_quad", "e:kvadrant", 15, "in"), ("match", "shoda", 7, "f"),
]


def sheet_kalkulace(wb, rows, name, period):
    ws = wb.create_sheet("Kalkulace (živé vzorce)")
    idx = {k: i + 1 for i, (k, *_ ) in enumerate(KCOLS)}
    C = lambda key, r: f"{get_column_letter(idx[key])}{r}"

    def q_expr(r):  # quadrant with hysteresis referencing previous row's QUAD cell
        M, Lo, prev = C("MECH", r), C("LOAD", r), C("QUAD", r - 1)
        lh = f'OR({Lo}>=25,AND(OR({prev}="Přetížení",{prev}="Kritická kombinace"),{Lo}>=18))'
        mh = f'OR({M}>=25,AND(OR({prev}="Tichý drift",{prev}="Kritická kombinace"),{M}>=18))'
        return (f'=IF(AND({lh},{mh}),"Kritická kombinace",'
                f'IF({lh},"Přetížení",IF({mh},"Tichý drift","Stabilní")))')

    def formula(key, r):
        c = lambda k: C(k, r)
        sev = f"MAX({c('sessionSpike')}-1,{c('ratio')}-1,0)"
        F = {
            "p_tavr": f"=ROUND(MIN(4,MAX(0,{c('tavr_z')}-0.2))*17)",
            "p_gct": f"=ROUND(MIN(4,MAX(0,{c('gct_z')}-0.2))*13)",
            "p_cad": f"=ROUND(MIN(4,MAX(0,-{c('cad_z')}-0.2))*10)",
            "p_vosc": f"=ROUND(MIN(4,MAX(0,{c('vosc_z')}-0.2))*10)",
            "p_bal": f"=ROUND(MIN(3,MAX(0,{c('bal_exc')}-0.4))*22)",
            "p_dec": f"=IF({c('dec_trend')}>0.15,ROUND(MIN(1.2,MAX(0,{c('dec_trend')}-0.1))*26),0)",
            "p_gaitcv": f"=IF({c('gaitcv_ratio')}>=1.5,ROUND(MIN(14,({c('gaitcv_ratio')}-1.5)*14)),0)",
            # v0.6 — single-session spike is the load spine (bands from BJSM 2025 HRRs)
            "p_sspike": (f"=IF(AND({c('valid')}=1,{c('sessionSpike')}>1.1),"
                         f"IF({c('sessionSpike')}>2,ROUND(MIN(16,({c('sessionSpike')}-2)*16)+14),"
                         f"IF({c('sessionSpike')}>1.3,ROUND(MIN(14,({c('sessionSpike')}-1.3)*20)),"
                         f"ROUND(MIN(6,({c('sessionSpike')}-1.1)*25)))),0)"),
            "p_slatent": f"=IF(AND({c('valid')}=1,{c('spikeLatent')}>0),ROUND(MIN(16,{c('spikeLatent')}*22)),0)",
            "p_pspike": f"=IF(AND({c('valid')}=1,{c('paceSpike')}>1.06),ROUND(MIN(10,({c('paceSpike')}-1.06)*40)),0)",
            # ACWR demoted to grade-C context
            "p_ewma": (f"=IF({c('valid')}=1,IF({c('ratio')}>1.5,ROUND(MIN(12,({c('ratio')}-1.5)*18)+2),"
                       f"IF({c('ratio')}<0.7,10,0)),0)"),
            "p_hi": f"=IF(AND({c('valid')}=1,{c('hiAcute')}>=60,{c('hiRatio')}>1.5),ROUND(MIN(20,({c('hiRatio')}-1.5)*20)),0)",
            "p_creep": f"=IF(AND({c('valid')}=1,{c('ratio')}<1.3,{c('loadCreep')}>=1.15),ROUND(MIN(10,({c('loadCreep')}-1.15)*30)),0)",
            "p_mono": f"=IF({c('monotony')}>2.4,ROUND(MIN(20,({c('monotony')}-2.4)*7)),0)",
            "p_desc": f"=IF({c('descentSpike')}>1.45,ROUND(MIN(14,({c('descentSpike')}-1.45)*15)),0)",
            "p_steep": f"=IF({c('steepSpike')}>1.5,ROUND(MIN(12,({c('steepSpike')}-1.5)*12)),0)",
            "p_aer": f"=IF({c('aer_mean')}>5.5,ROUND(MIN(12,({c('aer_mean')}-5.5)*3)),0)",
            "p_hrv": f"=IF({c('hrv_z')}<=-1,ROUND(MIN(22,-{c('hrv_z')}*10)),0)",
            "p_rhr": f"=IF({c('rhr_z')}>=1.2,ROUND(MIN(18,{c('rhr_z')}*8)),0)",
            "p_hrvcv": f"=IF({c('hrvcv_ratio')}>=1.4,ROUND(MIN(10,({c('hrvcv_ratio')}-1.4)*14)),0)",
            # v0.6 — load × capacity interaction
            "p_lcap": (f"=IF(AND({c('valid')}=1,{c('capacityDeficit')}>=0.3,{sev}>0.1),"
                       f"ROUND(MIN(16,{c('capacityDeficit')}*{sev}*34)),0)"),
            "p_tsb": (f"=IF(AND({c('valid')}=1,{c('chronic')}>0),"
                      f"IF(({c('tsbBalance')}/{c('chronic')})<=-0.12,ROUND(MIN(12,(-({c('tsbBalance')}/{c('chronic')})-0.12)*90)),0),0)"),
            "p_sleep": f"=IF({c('sleep_debt')}>=4,ROUND(MIN(16,{c('sleep_debt')}*2.5)),0)",
            "p_sreg": f"=IF({c('sleepreg_ratio')}>=1.5,ROUND(MIN(14,({c('sleepreg_ratio')}-1.5)*12)),0)",
            "p_seff": f"=IF(AND({c('sleepeff_now')}>0,{c('sleepeff_now')}<0.85),ROUND(MIN(16,(0.85-{c('sleepeff_now')})*60)),0)",
            # frailty (injury history) multiplies the objective axes
            "MECH": f"=MIN(100,ROUND(SUM({c('p_tavr')}:{c('p_gaitcv')})*{c('frailty')}))",
            "LOAD": f"=MIN(100,ROUND(SUM({c('p_ewma')}:{c('p_tsb')})*{c('frailty')}))",
            "SYMP": f"=MIN(100,SUM({c('p_sleep')}:{c('p_seff')}))",
            "OVERALL": f"=ROUND(MIN(100,{c('MECH')}*0.38+{c('LOAD')}*0.3+{c('SYMP')}*0.52))",
            "TIER": f'=IF({c("OVERALL")}>=70,"vysoké",IF({c("OVERALL")}>=40,"sledovat","nízké"))',
            "QUAD": q_expr(r),
            "match": f'=IF(AND({c("MECH")}={c("eng_mech")},{c("LOAD")}={c("eng_load")},{c("SYMP")}={c("eng_symp")},{c("QUAD")}={c("eng_quad")}),"✓","≠")',
        }
        return F[key]

    ws.append([f"Kalkulace enginu {E.ENGINE_VERSION} — živé vzorce", name])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Jak číst", "Modré = vstupy z enginu (rolling z/EWMA/TRIMP/spike — to je kód). Žluté = ŽIVÉ VZORCE (body signálů). "
               "Zelené = osy → celkem → kvadrant. 'e:' = co spočítal engine, 'shoda' = ✓ když vzorce sedí. Změň modrou buňku a vše se přepočítá."])
    ws.append(["v0.6", "Páteří zátěže je skok v jednom běhu (běh spike×) + doznívající paměť + tempo; ACWR je jen kontext (grade C). "
               "Frailty× = násobič z anamnézy zranění na osy MECH a ZÁTĚŽ. kap×zát = interakce zátěž × snížená regenerace."])
    ws.append(["overall", "MECH·0.38 + ZÁTĚŽ·0.30 + PŘÍZN·0.52 (každá 0–100) · kvadrant s hysterezí (práh 25, drží do 18) · odkazuje na řádek výše"])
    ws.append([f"Období {period}", "Historie = objektivní engine z hodinek; self-report (bolest/obtíže) a frailty se v čase nepřehrávají, proto jsou příznaky jen ze spánku."])
    ws.append([])
    gband = ws.max_row + 1
    spans = [("VSTUPY (z enginu)", "date", "sleepeff_now"), ("BODY — MECHANIKA", "p_tavr", "p_gaitcv"),
             ("BODY — ZÁTĚŽ", "p_ewma", "p_tsb"), ("BODY — PŘÍZNAKY", "p_sleep", "p_seff"),
             ("VÝSLEDEK (vzorce)", "MECH", "QUAD"), ("ENGINE REF", "eng_mech", "match")]
    for lab, k0, k1 in spans:
        cell = ws.cell(gband, idx[k0], lab)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = GRP
        cell.alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=gband, start_column=idx[k0], end_row=gband, end_column=idx[k1])
    hrow = gband + 1
    header(ws, [h for (_, h, *_ ) in KCOLS], [w for (_, _, w, _) in KCOLS], row=hrow)
    ws.freeze_panes = f"C{hrow + 1}"

    r = hrow
    for row in rows:
        r += 1
        for key, _h, _w, kind in KCOLS:
            cell = ws.cell(r, idx[key])
            if kind == "in":
                cell.value = row.get(key)
                if key not in ("date",) and not key.startswith("eng"):
                    cell.fill = IN_FILL
                elif key.startswith("eng"):
                    cell.fill = PatternFill("solid", fgColor="EEECEC")
            else:
                cell.value = formula(key, r)
                cell.fill = RES_FILL if key in ("MECH", "LOAD", "SYMP", "OVERALL", "TIER", "QUAD") else F_FILL
            cell.alignment = Alignment(horizontal="center")
        ws.cell(r, idx["QUAD"]).fill = PatternFill("solid", fgColor=QFILL.get(row["_q"], "E7F5EC"))


# ---- backtest timeline sheets
BT = [
    ("date", "Datum", 11), ("runs", "Běhů", 6), ("days", "Dnů hist.", 8), ("conf", "Konfid.", 8),
    ("runKm7", "Běh 7d km", 9), ("acute", "Akut AU", 8), ("sessionSpike", "Běh spike×", 9),
    ("sessionSpikeKm", "Nejdel. km", 9), ("spikeLatent", "Doznívá", 8), ("paceSpike", "Tempo×", 7),
    ("safeLongRunKm", "Strop LR km", 10), ("ratio", "ACWR", 7), ("hiRatio", "HI ×", 7),
    ("loadCreep", "creep×", 7), ("crossLoad7", "Jiný sport AU", 11), ("monotony", "Mono", 7),
    ("capacityDeficit", "Kap.def", 8),
    ("tavr_z", "vert.p. z", 8), ("gct_z", "GCT z", 7), ("cad_z", "kad. z", 7), ("vosc_z", "osc. z", 7),
    ("duty_z", "duty z", 7), ("gaitcv_ratio", "gaitCV×", 8),
    ("regen", "Regen.", 7), ("regen_d", "Δ noc", 7), ("sleepeff_now", "sp.efek", 7),
    ("eng_mech", "MECH", 7), ("eng_load", "ZÁTĚŽ", 7), ("eng_symp", "PŘÍZN.", 7), ("eng_overall", "CELKEM", 8),
    ("eng_quad", "Kvadrant", 15), ("trans", "Přechod", 9), ("signals", "Hlavní signály", 60),
]


def sheet_backtest(wb, title, rows, name, period, daily=False):
    ws = wb.create_sheet(title)
    cad = "den po dni" if daily else "týden po týdnu"
    for rr in [
        [f"Engine {E.ENGINE_VERSION} — historický replay ({cad})", name],
        ["Období dat", period],
        ["Pozn.", "Objektivní engine z hodinek do daného data. Zátěž v jednotkách (AU) přes všechny sporty; páteří je skok v jednom běhu; mechanika terénně očištěná."],
        [],
    ]:
        ws.append(rr)
    ws["A1"].font = Font(bold=True, size=13)
    ws.cell(3, 1).font = Font(bold=True)
    hrow = ws.max_row + 1
    header(ws, [h for _, h, _ in BT], [w for _, _, w in BT], row=hrow)
    keys = [k for k, _, _ in BT]
    qi = keys.index("eng_quad") + 1
    for row in rows:
        ws.append([round(row[k], 2) if isinstance(row.get(k), float) else row.get(k) for k in keys])
        ws.cell(ws.max_row, qi).fill = PatternFill("solid", fgColor=QFILL.get(row["_q"], "FFFFFF"))
        ws.cell(ws.max_row, keys.index("signals") + 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.auto_filter.ref = f"A{hrow}:{get_column_letter(len(BT))}{ws.max_row}"


def sheet_signals(wb):
    ws = wb.create_sheet("Signály")
    ws.append([f"Signály enginu {E.ENGINE_VERSION}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Grade:", " · ".join(f"{k}: {v}" for k, v in SD.GRADE_NOTE.items())])
    ws.append([])
    r0 = ws.max_row + 1
    header(ws, ["id", "signál", "grade", "vzorec (fx)", "proč", "limit", "kdy zmizí"], [11, 30, 7, 44, 54, 48, 38], row=r0)
    for sid, doc in SD.SIG_DOC.items():
        ws.append([sid, doc.get("t", ""), doc.get("g", ""), doc.get("fx") or "", doc.get("why", ""), doc.get("limit", ""), doc.get("clear", "")])
    for r in range(r0 + 1, ws.max_row + 1):
        for c in range(1, 8):
            ws.cell(r, c).alignment = Alignment(vertical="top", wrap_text=True)


def main():
    weekly, name, period = replay(7)
    daily, _, _ = replay(1)
    bad = 0
    prev = None
    for row in weekly:
        M, Lo, Sy, ov, q = _calc(row, prev)
        prev = row["_q"]
        if (M, Lo, Sy) != (row["eng_mech"], row["eng_load"], row["eng_symp"]) or q != row["_q"]:
            bad += 1
            print(f"  MISMATCH {row['date']}: calc M{M} L{Lo} S{Sy} {q} vs engine M{row['eng_mech']} L{row['eng_load']} S{row['eng_symp']} {row['_q']}")
    print(f"validace vzorců vs engine: {len(weekly) - bad}/{len(weekly)} sedí")

    wb = Workbook()
    wb.remove(wb.active)
    sheet_kalkulace(wb, weekly, name, period)
    sheet_backtest(wb, "Backtest — týdně", weekly, name, period, daily=False)
    sheet_backtest(wb, "Backtest — denně", daily, name, period, daily=True)
    sheet_signals(wb)
    wb.save(OUT)
    print(f"Hotovo → {OUT}\nListy: {wb.sheetnames} · týdnů {len(weekly)} · dnů {len(daily)}")


if __name__ == "__main__":
    main()
