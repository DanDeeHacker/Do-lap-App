"""Backtest workbook for the *current* engine (v0.4.x), same spirit as
dosslap_engine_v04_backtest.xlsx but reflecting every improvement since:
load axis in training-load units (AU, incl. cross-training), calendar weeks,
terrain-cleaned cadence / vertical oscillation, single-night recovery + Δ,
continuous mechanical drift, new signals (cad, vosc, fatigue, pain recurrence).

Weekly historical replay of assess() over run-0012's real Garmin history.
Self-reported inputs (check-in / diary / injury) are NOT replayed — the
historical timeline is the objective wearable-derived engine only.
"""
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app import models
from app.db import Base
from app.metrics import engine as E
from app.metrics import sig_doc as SD

DB = "dosslap.db"
RID = "run-0012"
OUT = "/Users/danieltrnovec/Downloads/dosslap/dosslap_engine_v048_backtest.xlsx"

HEAD = PatternFill("solid", fgColor="0A2540")
QFILL = {"stable": "C6EFCE", "overreaching": "FFEB9C", "silent": "BDD7EE", "critical": "FFC7CE"}
TIER_CZ = {"ok": "OK", "watch": "Sledovat", "alert": "Vysoké"}


def header(ws, names, widths, row=1):
    ws.append(names)
    for i, n in enumerate(names, 1):
        c = ws.cell(row=row, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = f"A{row + 1}"
    ws.row_dimensions[row].height = 30


def _rows_of(session, model, rid):
    cols = [c.name for c in model.__table__.columns]
    return [{c: getattr(r, c) for c in cols} for r in session.query(model).filter(model.runner_id == rid).all()]


def replay(step_days=7):
    src = sessionmaker(bind=create_engine(f"sqlite:///{DB}"))()
    runner = src.query(models.Runner).filter(models.Runner.id == RID).first()
    rdata = {c.name: getattr(runner, c.name) for c in models.Runner.__table__.columns}
    acts = _rows_of(src, models.Activity, RID)
    daily = _rows_of(src, models.DailyMetric, RID)
    src.close()

    first = min(a["started_at"][:10] for a in acts)
    last = max(a["started_at"][:10] for a in acts)
    ad = date.fromisoformat(first) + timedelta(days=42)
    end = date.fromisoformat(last)
    asofs = []
    while ad <= end:
        asofs.append(ad)
        ad += timedelta(days=step_days)
    if not asofs or asofs[-1] != end:
        asofs.append(end)

    orig = E.today_date
    out = []
    try:
        for ad in asofs:
            E.today_date = (lambda d: (lambda: d))(ad)
            eng = create_engine("sqlite://")
            Base.metadata.create_all(eng)
            ts = sessionmaker(bind=eng)()
            ts.add(models.Runner(**rdata))
            cut = ad.isoformat()
            for a in acts:
                if a["started_at"][:10] <= cut:
                    ts.add(models.Activity(**{k: v for k, v in a.items() if k != "id"}))
            for d in daily:
                if d["date"][:10] <= cut:
                    ts.add(models.DailyMetric(**{k: v for k, v in d.items() if k != "id"}))
            ts.commit()
            if out:  # seed previous quadrant for hysteresis
                ts.add(models.Assessment(runner_id=RID, quadrant=out[-1]["_q"], tier="ok",
                                         mech=0, load=0, symp=0, overall=0, engine_version=E.ENGINE_VERSION))
                ts.commit()
            a = E.assess(ts, RID)
            conf = a["confidence"] or {}
            L = a.get("loadDetail") or {}
            g = lambda k, sub="z": (a.get(k) or {}).get(sub)
            rcv = a.get("rcv") or {}
            out.append({
                "as_of": cut, "runs": sum(1 for x in acts if x["started_at"][:10] <= cut),
                "days": conf.get("days"), "base": conf.get("baseSessions"), "conf": conf.get("value"),
                "run7": L.get("runKm7"), "acute": L.get("acute"), "chronic": L.get("chronic"),
                "ewma": L.get("ratio") if L.get("valid") else None, "mono": L.get("monotony"),
                "cross7": L.get("crossLoad7"),
                "tavr_z": g("tavr"), "gct_z": g("gct"), "cad_z": g("cadence"), "vosc_z": g("vosc"),
                "regen": rcv.get("score"), "regen_d": rcv.get("scoreDelta"),
                "hrv_z": (rcv.get("hrv") or {}).get("z"), "rhr_z": (rcv.get("rhr") or {}).get("z"),
                "sleep_debt": (rcv.get("sleep") or {}).get("debt"),
                "mech": a["mech"], "load": a["load"], "symp": a["symp"], "overall": a["overall"],
                "tier": TIER_CZ.get(a["tier"], a["tier"]), "quad": SD.QUAD[a["quadrant"]]["t"], "_q": a["quadrant"],
                "signals": " · ".join(f"{s['name']} (+{s['pts']}, {s['grade']})" for s in (a.get("signals") or [])[:4]) or "—",
            })
            ts.close()
            eng.dispose()
    finally:
        E.today_date = orig
    # mark quadrant transitions
    for i, r in enumerate(out):
        r["trans"] = "← ZMĚNA" if i and r["_q"] != out[i - 1]["_q"] else None
    return out, runner.name, f"{first} → {last}"


COLS = [
    ("as_of", "Datum", 12), ("runs", "Běhů celkem", 10), ("days", "Dnů historie", 10),
    ("base", "Baseline běhů", 11), ("conf", "Konfidence", 10),
    ("run7", "Běh 7d km", 9), ("acute", "Akut zátěž (AU)", 12), ("chronic", "Chron zátěž (AU)", 12),
    ("ewma", "EWMA poměr", 10), ("mono", "Monotón.", 9), ("cross7", "Jiný sport 7d (AU)", 13),
    ("tavr_z", "Vert.poměr z", 11), ("gct_z", "GCT z", 8), ("cad_z", "Kadence z", 10), ("vosc_z", "Vert.oscil. z", 11),
    ("regen", "Regenerace", 10), ("regen_d", "Δ přes noc", 9), ("hrv_z", "HRV z", 8), ("rhr_z", "Klid.tep z", 10),
    ("sleep_debt", "Spánek dluh", 10),
    ("mech", "Mechanika", 10), ("load", "Zátěž", 8), ("symp", "Příznaky", 9), ("overall", "Celkem", 8),
    ("tier", "Úroveň", 9), ("quad", "Kvadrant", 13), ("trans", "Přechod", 10), ("signals", "Hlavní signály", 62),
]


def sheet_backtest(wb, title, rows, name, period, daily=False):
    ws = wb.create_sheet(title)
    cadence = "den po dni" if daily else "týden po týdnu"
    for r in [
        [f"Engine {E.ENGINE_VERSION} — historický replay ({cadence})", name],
        ["Období dat", period],
        ["Osy", "overall = mechanika·0.38 + zátěž·0.30 + příznaky·0.52 (každá 0–100). Kvadrant = mechanika × zátěž (práh 25, hystereze 18)."],
        ["Zátěž", "Nově v jednotkách tréninkové zátěže (AU, TRIMP z tepu) přes VŠECHNY sporty — ne jen běžecké km. 'Běh 7d km' je jen objem běhu."],
        ["Mechanika", "z = odchylka od vlastní normy ve srovnatelném terénu×tempu. Skóre driftu je spojité — i jemný drift (z>0,2) zvedá skóre z nuly."
                      + (" Mechanika se mění jen v dny s během; regenerace a zátěž se hýbou i mezi běhy." if daily else "")],
        ["Pozn.", "Každý řádek = stav enginu jen z dat do daného data. Self-report (check-in / deník / zranění) se v historii nepřehrává — tady je jen objektivní engine z hodinek."],
        [],
    ]:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=13)
    for i in (3, 4, 5, 6, 7):
        ws.cell(i, 1).font = Font(bold=True)
    r0 = ws.max_row + 1
    header(ws, [c[1] for c in COLS], [c[2] for c in COLS], row=r0)
    keys = [c[0] for c in COLS]
    qi = keys.index("quad") + 1
    for r in rows:
        ws.append([round(r[k], 2) if isinstance(r.get(k), float) else r.get(k) for k in keys])
        ws.cell(ws.max_row, qi).fill = PatternFill("solid", fgColor=QFILL.get(r["_q"], "FFFFFF"))
        ws.cell(ws.max_row, keys.index("signals") + 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(COLS))}{ws.max_row}"


LEGENDA = [
    ("Konfidence", "Spolehlivost baseline 0–1. Pod 0,6 se mechanické signály (vert.poměr, GCT, kadence, oscilace) vůbec nepočítají."),
    ("Akut / Chron zátěž (AU)", "EWMA tréninkové zátěže v jednotkách (TRIMP z tepu) přes všechny sporty. Nahradilo km — kolo/plavání/silovka se teď počítají."),
    ("EWMA poměr", "Akutní vs chronická zátěž (AU). Pásmo ~0,7–1,3 klidné; 1,3–1,5 mírně zvýšené; >1,5 zřetelně zvýšené riziko."),
    ("Jiný sport 7d (AU)", "Kolik z týdenní zátěže přišlo z neběžeckých sportů (běžecké km jsou zvlášť). Vstupuje do poměru i monotónnosti."),
    ("Vert.poměr / GCT / Kadence / Vert.oscil. z", "Odchylka běžecké mechaniky od vlastní normy ve srovnatelném terénu×tempu. Bad-směr: vert.poměr↑, GCT↑, kadence↓, oscilace↑. Skóre driftu je spojité."),
    ("Regenerace / Δ přes noc", "Jednonoční skóre 0–100 (HRV/klid.tep/spánek té noci vs baseline) a jeho změna oproti předchozí noci."),
    ("HRV z / Klid.tep z / Spánek dluh", "Regenerace z hodinek proti baseline. HRV z ≤ −1 nebo klid.tep z ≥ 1,2 = potlačená regenerace; spánek dluh h/týden."),
    ("Mechanika / Zátěž / Příznaky", "Tři osy skóre 0–100, nesčítají se do jednoho. Kvadrant je jejich kombinace, celkem je vážený součet."),
    ("Kvadrant", "Stabilní | Přetížení (zátěž↑, technika drží) | Tichý drift (technika driftuje bez objemu) | Kritické (obojí)."),
    ("Přechod", "Řádek, kde se kvadrant změnil proti předchozímu týdnu — moment, který stojí za pohled."),
    ("Hlavní signály", "Top 4 aktivní signály (+body, evidence grade A/B/C) v daném týdnu."),
]


def sheet_legenda(wb):
    ws = wb.create_sheet("Legenda")
    header(ws, ["Sloupec", "Význam"], [30, 120])
    for row in LEGENDA:
        ws.append(list(row))
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 1).font = Font(bold=True)


# axis · trigger · points — mirrors current assess()
SIG_META = {
    "tavr": ("mechanika", "z > 0.2 (signál od 0.6)", "clamp(z−0.2,0,4)·17 — spojité"),
    "gct": ("mechanika", "z > 0.2 (signál od 0.6)", "clamp(z−0.2,0,4)·13 — spojité"),
    "cad": ("mechanika", "−z > 0.2 (kadence klesá)", "clamp(−z−0.2,0,4)·10 — spojité"),
    "vosc": ("mechanika", "z > 0.2 (oscilace roste)", "clamp(z−0.2,0,4)·10 — spojité"),
    "bal": ("mechanika", "excursion > 0.4 (signál od 0.8)", "clamp(exc−0.4,0,3)·22"),
    "dec": ("mechanika", "trend > 0.15 (signál od 0.35)", "clamp(trend−0.1,0,1.2)·26"),
    "ewma": ("zátěž", "ratio > 1.5 (nebo < 0.7)", "clamp((ratio−1.5)·70,0,42)+10 · <0.7 ⇒ 22"),
    "ewma_mild": ("zátěž", "1.3 < ratio ≤ 1.5", "clamp((ratio−1.3)·35,0,14)"),
    "mono": ("zátěž", "monotony > 2.4", "clamp((mono−2.4)·7,0,26)"),
    "desc": ("zátěž", "descentSpike > 1.45", "clamp((spike−1.45)·30,0,26)"),
    "desc_steep": ("zátěž", "steepSpike > 1.5", "clamp((spike−1.5)·20,0,22)"),
    "aer": ("zátěž", "aer.mean > 5.5 %", "clamp((mean−5.5)·3,0,12)"),
    "hrv": ("zátěž", "hrv.z ≤ −1.0", "clamp(−z·10,0,22)"),
    "rhr": ("zátěž", "rhr.z ≥ 1.2", "clamp(z·8,0,18)"),
    "hrvcv": ("zátěž", "CV ratio ≥ 1.4", "clamp((ratio−1.4)·20,0,16)"),
    "tsb": ("zátěž", "tsbBalance/chronic ≤ −0.12", "clamp((−rel−0.12)·120,0,18)"),
    "taper": ("zátěž", "≤21 dní do závodu A (load≥25 || ratio>1.3)", "clamp((21−d)/21·14,4,14)"),
    "niggle": ("příznaky", "niggleCount ≥ 3 / 21 dní", "clamp(niggleCount·9,0,40)"),
    "pain": ("příznaky", "denní check-in bolest ≥1/≥3/≥6 (okno 4 dny)", "8/26/44 + bonus: back-to-back(2d) až +26 · opakování(28d) až +18"),
    "sore": ("příznaky", "soreness ≥ 7", "10"),
    "fatigue": ("příznaky", "check-in únava ≥ 6", "clamp((stress−5)·2.5,0,12)"),
    "sleep": ("příznaky", "debt ≥ 4 h/týd", "clamp(debt·2.5,0,16)"),
    "sleepreg": ("příznaky", "SD ratio ≥ 1.5", "clamp((ratio−1.5)·12,0,14)"),
    "feel": ("příznaky", "n≥6 A feelingTrend ≤ −0.12", "10"),
    "stiffness": ("příznaky", "ignoreRate ≥0.5 (≥3×) / trend ≥0.15", "clamp(ignoreRate·20,0,18) / 10"),
    "injury": ("příznaky", "aktivní OSTRC report (28 dní)", "potvrzeno clamp(sev·0.5,0,46) [A] · self clamp(sev·0.34,0,32) [B]"),
    "hist": ("příznaky", "prior_injury ≤ 12 měs.", "16"),
}
SIG_ORDER = ["tavr", "gct", "cad", "vosc", "bal", "dec", "ewma", "ewma_mild", "mono", "desc",
             "desc_steep", "aer", "hrv", "rhr", "hrvcv", "tsb", "taper", "niggle", "pain",
             "sore", "fatigue", "sleep", "sleepreg", "feel", "stiffness", "injury", "hist"]


def sheet_signals(wb):
    ws = wb.create_sheet("Signály")
    ws.append([f"Signály enginu {E.ENGINE_VERSION} — osa, práh, body, grade, vzorec"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Tier:", "overall ≥ 70 = vysoké riziko · ≥ 40 = sledovat · jinak nízké"])
    ws.append(["Grade:", " · ".join(f"{k}: {v}" for k, v in SD.GRADE_NOTE.items())])
    ws.append([])
    r0 = ws.max_row + 1
    header(ws, ["id", "signál", "osa", "grade", "práh (trigger)", "body (formula)",
                "proč", "kde selhává (limit)", "kdy zmizí (clear)"],
           [11, 32, 11, 7, 34, 40, 58, 52, 40], row=r0)
    for sid in SIG_ORDER:
        doc = SD.SIG_DOC.get(sid, {})
        axis, trig, pts = SIG_META.get(sid, ("", "", ""))
        ws.append([sid, doc.get("t", ""), axis, doc.get("g", ""), trig, pts,
                   doc.get("why", ""), doc.get("limit", ""), doc.get("clear", "")])
    for r in range(r0 + 1, ws.max_row + 1):
        for c in range(1, 10):
            ws.cell(r, c).alignment = Alignment(vertical="top", wrap_text=True)


DOC = [
    ("PŘEHLED", ""),
    ("Co engine dělá", "Z hodinek (běhy + denní wellness) a self-reportu počítá tři nezávislé osy rizika 0–100 a z nich kvadrant. Není to diagnóza — podporuje rozhodnutí."),
    ("Tři osy", "Mechanika (jak se mění běžecká technika), Zátěž (kolik toho tělo unese), Příznaky (co běžec cítí/hlásí). Nesčítají se do jednoho čísla."),
    ("Celkem (overall)", "overall = mechanika·0.38 + zátěž·0.30 + příznaky·0.52 (clamp 0–100). Příznaky mají největší váhu — reálná bolest/hlášení váží nejvíc."),
    ("Kvadrant", "Kombinace os s hysterezí: osa 'hoří' od skóre 25, 'chladne' až pod 18. Stabilní / Přetížení (zátěž↑) / Tichý drift (mechanika bez objemu) / Kritické (obojí)."),
    ("Úroveň (tier)", "overall ≥ 70 = vysoké riziko · ≥ 40 = sledovat · jinak nízké."),
    ("PŘÍJEM DAT", ""),
    ("Sporty", "Aktivity se tagují sportem. Jen běh jde do mechaniky (kadence, vert.poměr, GCT, terénní buckety). Všechny sporty jdou do systémové zátěže."),
    ("Čištění senzorů", "0 / záporná hodnota u běžecké dynamiky nebo tepu = senzor to nezachytil → bere se jako chybějící (None), ne jako reálná 0. Jinak by nula stáhla průměr metriky."),
    ("Terénní profil (bucket)", "povrch × sklon × tempo. Sklon: klesání / stoupání (čisté převýšení >12 m/km), kopcovitě (hodně nahoru i dolů ≥30 m/km, typicky trail), rovina. Mechanika se porovnává jen ve stejném profilu."),
    ("ZÁTĚŽ", ""),
    ("Jednotka (AU)", "Tréninková zátěž v jednotkách = Banister TRIMP z tepové odezvy, napříč sporty (běh i kolo/plavání/silovka na jedné škále). Fallback: Garmin training load, jinak délka×intenzita."),
    ("Akutní / chronická", "EWMA denní zátěže (7 vs 28 dní) × 7. Poměr = akut/chron (ACWR): ~0,7–1,3 klidné, 1,3–1,5 mírně zvýšené, >1,5 zřetelně zvýšené."),
    ("Monotónnost / strain", "Foster: průměr/SD denní zátěže. Vysoká monotónnost (chybí lehké dny) zvyšuje riziko."),
    ("Objem vs zátěž", "Kilometry (objem běhu) jsou po KALENDÁŘNÍCH týdnech (Po–Ne) + trailing 7 dní. Zátěžová osa jede z AU, ne z km."),
    ("MECHANIKA", ""),
    ("Terénní očištění", "Vert.poměr, GCT, kadence, délka kroku, oscilace = z-odchylka od vlastní normy ve shodném profilu terénu×tempa (drift-z po bucketech, per-bucket z zastropované ±4 proti outlierům)."),
    ("Spojité skóre", "Každá metrika přispívá úměrně tomu, jak moc driftuje za malou mrtvou zónou (z>0,2) v 'bad' směru (vert.poměr↑, GCT↑, kadence↓, oscilace↑). I jemný drift zvedá skóre z nuly; signál se vypíše od z≈0,6."),
    ("Konfidence", "Spolehlivost baseline 0–1 (shodné běhy / dny historie / počet baseline běhů). Pod 0,6 se mechanické signály vůbec nepočítají."),
    ("REGENERACE", ""),
    ("Jednonoční skóre", "0–100 (50 = vlastní baseline) z HRV / klid.tepu / spánku TÉ noci z-skórovaných proti 28denní baseline. Δ = rozdíl oproti předchozí noci (reálný noční pohyb, ne 7denní průměr)."),
    ("PŘÍZNAKY (self-report)", ""),
    ("Denní check-in", "Bolest (8/26/44), svalová ztuhlost (≥7 → 10), únava (≥6 → až 12). Okno 4 dny — je to 'dnešní' signál. Nálada se ukládá, ale NEskóruje (kontrola korelace s rizikem)."),
    ("Místo bolesti", "Počítá se přes opakování: stejné běžecky-relevantní místo znovu do 2 dnů (neustupuje mezi běhy) = akutní přetížení (silný bonus); ≥3× za 28 dní = chronické opakování (mírný bonus). Jednorázové/neběžecké → jen základ."),
    ("Varování bolesti", "Bolest > 3/10 → aplikace upozorní běžce (banner na Dnes + hláška v check-inu), aby zvážil odpočinek/konzultaci."),
    ("Týdenní check-in", "OSTRC-H (4 otázky × 0/8/17/25 → závažnost 0–100), okno 28 dní. Self-report grade B, potvrzeno fyziem grade A. 'Bez obtíží' se ukládá jako negativní datapoint."),
    ("Deník běhů", "Opakované bolestivé místo z hodnocení běhů (niggle) — frekvence bolesti po trénincích."),
]


def sheet_dokumentace(wb, name, period, n_week, n_day):
    ws = wb.create_sheet("Dokumentace")
    ws.append([f"Engine {E.ENGINE_VERSION} — jak model funguje", name])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Data", f"{period} · běžec run-0012 (reálný Garmin) · {n_week} týdenních a {n_day} denních snímků"])
    ws.append(["Viz také", "List 'Signály' = všechny signály s prahy/vzorci/grade. Listy 'Backtest' = replay v čase."])
    ws.append([])
    header(ws, ["Téma", "Popis"], [26, 130], row=ws.max_row + 1)
    for k, v in DOC:
        ws.append([k, v])
        if v == "":  # section header
            c = ws.cell(ws.max_row, 1)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEAD
            ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=2)
        else:
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            ws.cell(ws.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")


def main():
    wb = Workbook()
    wb.remove(wb.active)
    weekly, name, period = replay(7)
    daily, _, _ = replay(1)
    sheet_backtest(wb, "Backtest — týdně", weekly, name, period, daily=False)
    sheet_backtest(wb, "Backtest — denně", daily, name, period, daily=True)
    sheet_dokumentace(wb, name, period, len(weekly), len(daily))
    sheet_signals(wb)
    sheet_legenda(wb)
    wb.save(OUT)
    print(f"Hotovo → {OUT}")
    print("Listy:", wb.sheetnames, "· týdnů:", len(weekly), "· dnů:", len(daily))


if __name__ == "__main__":
    main()
